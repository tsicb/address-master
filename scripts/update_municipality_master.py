#!/usr/bin/env python3
import json
import re
import sys
import urllib.request
from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

SOURCE_PAGE = "https://www.soumu.go.jp/denshijiti/code.html"
FALLBACK_XLSX = "https://www.soumu.go.jp/main_content/000925835.xlsx"
OUT = Path("municipality_master.json")

PREF_SUFFIXES = ("都", "道", "府", "県")

def fetch_bytes(url, timeout=30):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 address-master-updater"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read()

def discover_xlsx():
    try:
        html = fetch_bytes(SOURCE_PAGE).decode("utf-8", errors="ignore")
        pos = html.find("都道府県コード及び市区町村コード")
        scope = html[pos:pos+12000] if pos >= 0 else html
        urls = re.findall(r'https?://www\.soumu\.go\.jp/main_content/\d+\.xlsx', scope)
        if not urls:
            rels = re.findall(r'["\'](/main_content/\d+\.xlsx)["\']', scope)
            urls = ["https://www.soumu.go.jp" + x for x in rels]
        if urls:
            return urls[0]
    except Exception as e:
        print(f"warning: source-page discovery failed: {e}", file=sys.stderr)
    return FALLBACK_XLSX

def sval(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()

def is_code(v):
    s = re.sub(r"\D", "", sval(v))
    return s if len(s) in (5, 6) else ""

def find_header(ws):
    best = None
    for r in range(1, min(ws.max_row, 20) + 1):
        vals = [sval(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 12) + 1)]
        score = sum(1 for x in vals if any(k in x for k in ("団体コード","都道府県名","市区町村名","政令指定都市","区名")))
        if best is None or score > best[0]:
            best = (score, r, vals)
    return best[1], best[2]

def col_index(headers, predicates):
    for i, h in enumerate(headers, start=1):
        if all(p in h for p in predicates) and "カナ" not in h and "かな" not in h:
            return i
    return None

def extract_rows(ws):
    header_row, headers = find_header(ws)
    code_col = col_index(headers, ("団体コード",)) or 1
    pref_col = col_index(headers, ("都道府県名",))
    city_col = col_index(headers, ("市区町村名",))
    designated_col = col_index(headers, ("政令指定都市",))
    ward_col = col_index(headers, ("区名",))

    records = []
    for r in range(header_row + 1, ws.max_row + 1):
        code_raw = is_code(ws.cell(r, code_col).value)
        if not code_raw:
            continue
        std = code_raw[:5].zfill(5)
        pref = sval(ws.cell(r, pref_col).value) if pref_col else ""
        city = sval(ws.cell(r, city_col).value) if city_col else ""
        designated = sval(ws.cell(r, designated_col).value) if designated_col else ""
        ward = sval(ws.cell(r, ward_col).value) if ward_col else ""

        if designated and ward:
            city = designated + ward
        elif designated and not city:
            city = designated
        elif ward and not city:
            city = ward

        # Fallback for layouts where the first text columns follow the code.
        if not pref or not city:
            texts=[]
            for c in range(code_col+1, min(ws.max_column, code_col+7)+1):
                v=sval(ws.cell(r,c).value)
                if v and not re.fullmatch(r"[\d\-./]+", v):
                    texts.append(v)
            kanji=[x for x in texts if re.search(r"[一-龯ぁ-んァ-ヶ]", x) and not re.fullmatch(r"[ｦ-ﾟ ]+",x)]
            if not pref:
                pref=next((x for x in kanji if x.endswith(PREF_SUFFIXES)), "")
            if not city:
                candidates=[x for x in kanji if x!=pref and not x.endswith(PREF_SUFFIXES)]
                if candidates:
                    city="".join(candidates[:2]) if len(candidates)>=2 and candidates[0].endswith("市") and candidates[1].endswith("区") else candidates[0]

        if pref and city:
            records.append((std, pref, city))
    return records

def main():
    url = discover_xlsx()
    print("source:", url)
    try:
        xlsx = fetch_bytes(url)
    except Exception:
        if url != FALLBACK_XLSX:
            print("warning: discovered xlsx failed, using fallback", file=sys.stderr)
            url = FALLBACK_XLSX
            xlsx = fetch_bytes(url)
        else:
            raise

    wb = load_workbook(BytesIO(xlsx), read_only=True, data_only=True)
    raw=[]
    for ws in wb.worksheets:
        rows=extract_rows(ws)
        print(f"sheet {ws.title}: {len(rows)} rows")
        raw.extend(rows)

    # Deduplicate by standard code; prefer longer city name when a duplicate exists.
    by_code={}
    for std,pref,city in raw:
        cur=by_code.get(std)
        if cur is None or len(city)>len(cur[1]):
            by_code[std]=(pref,city)

    # Remove designated-city parent rows when ward rows exist (e.g. 01100 札幌市).
    remove=set()
    items=list(by_code.items())
    for std,(pref,city) in items:
        if not city.endswith("市"):
            continue
        for std2,(pref2,city2) in items:
            if std2!=std and pref2==pref and city2.startswith(city) and city2.endswith("区"):
                remove.add(std)
                break
    for code in remove:
        by_code.pop(code,None)

    prefectures={}
    municipalities={}
    for std,(pref,city) in sorted(by_code.items()):
        p2=std[:2]
        prefectures[p2]=pref
        municipalities[std]=city

    if len(prefectures)!=47:
        raise RuntimeError(f"expected 47 prefectures, got {len(prefectures)}")
    if not (1850 <= len(municipalities) <= 1950):
        raise RuntimeError(f"unexpected municipality count: {len(municipalities)}")
    for code in ("01101","12227","13101"):
        if code not in municipalities:
            raise RuntimeError(f"required code missing: {code}")

    payload={
        "version": date.today().isoformat(),
        "source": url,
        "prefectures": dict(sorted(prefectures.items())),
        "municipalities": dict(sorted(municipalities.items()))
    }
    OUT.write_text(json.dumps(payload,ensure_ascii=False,separators=(",",":"))+"\n",encoding="utf-8")
    print(f"wrote {OUT}: prefectures={len(prefectures)} municipalities={len(municipalities)}")

if __name__=="__main__":
    main()
